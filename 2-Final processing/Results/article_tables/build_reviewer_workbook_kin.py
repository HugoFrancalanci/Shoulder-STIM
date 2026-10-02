# =========================================================================
# build_reviewer_workbook_kin.py
# =========================================================================
# Author     :   H. Francalanci
#                Biomechanics and Translational Research in Surgery Group
#                University of Geneva
# License    :   Creative Commons Attribution-NonCommercial 4.0 International License
# Date       :   October 2026
# -------------------------------------------------------------------------
# Description :  Builds the reviewer workbook (kinematics : GH + ST, SPM1D group/individual, curves, non-interpretable zone) from the JSON export
#                written by generate_reviewer_workbook_kin.m — never run directly, the
#                MATLAB script calls it :  python build_reviewer_workbook_kin.py <export.json> <out.xlsx>
#                Formatted sheets (Arial), live Excel formulas (Holm-
#                Bonferroni decision recomputed from the p-values, group
#                mean/SD curves, durations...), consistency checks against
#                the MATLAB decisions. fullCalcOnLoad is set so Excel
#                computes every formula when the file is opened.
# Dependencies : Python 3, openpyxl
# =========================================================================
import json, math, sys
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from openpyxl.comments import Comment

SRC = sys.argv[1]
OUT = sys.argv[2]
d = json.load(open(SRC, encoding='utf-8'))

FONT = 'Arial'
F_BASE = Font(name=FONT, size=10)
F_HEAD = Font(name=FONT, size=10, bold=True, color='FFFFFF')
F_TITLE = Font(name=FONT, size=14, bold=True)
F_BOLD = Font(name=FONT, size=10, bold=True)
F_NOTE = Font(name=FONT, size=9, italic=True, color='555555')
FILL_HEAD = PatternFill('solid', start_color='3F4E66')
FILL_SIG = PatternFill('solid', start_color='E8F1E4')
FILL_ZONE = PatternFill('solid', start_color='EDEDED')
THIN = Side(style='thin', color='BFBFBF')
BORDER = Border(bottom=THIN)
P_FMT = '[<0.001]"<0.001";0.000'
P_FMT4 = '[<0.0001]"<0.0001";0.0000'
ANG = '0.0'
PCT = '0.0'

def val(x):
    if x is None: return None
    if isinstance(x, float) and math.isnan(x): return None
    return x

def header(ws, row, cols, widths=None, freeze=True):
    for j, c in enumerate(cols, 1):
        cell = ws.cell(row=row, column=j, value=c)
        cell.font = F_HEAD; cell.fill = FILL_HEAD
        cell.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
    ws.row_dimensions[row].height = 42
    if widths:
        for j, w in enumerate(widths, 1):
            ws.column_dimensions[get_column_letter(j)].width = w
    if freeze:
        ws.freeze_panes = ws.cell(row=row+1, column=1)
        ws.auto_filter.ref = f"A{row}:{get_column_letter(len(cols))}{row}"

def style_body(ws, r1, r2, ncol):
    for r in range(r1, r2+1):
        for c in range(1, ncol+1):
            cell = ws.cell(row=r, column=c)
            if cell.font is None or not cell.font.bold:
                cell.font = F_BASE
            cell.border = BORDER

def title(ws, text, sub):
    ws['A1'] = text; ws['A1'].font = F_TITLE
    ws['A2'] = sub; ws['A2'].font = F_NOTE

CONV = {
    ('Glenohumeral', 'Elevation'): 'negative = elevation (flexion); positive = toward extension',
    ('Glenohumeral', 'External (-) / internal (+) rotation'): 'negative = external rotation; positive = internal rotation',
    ('Glenohumeral', 'Plane of elevation'): 'negative = anterior; positive = posterior plane of elevation',
    ('Scapulothoracic', 'Lateral (-) / medial (+) rotation'): 'negative = lateral (upward) rotation; positive = medial rotation',
    ('Scapulothoracic', 'Protraction (+) / retraction (-)'): 'positive = protraction; negative = retraction',
    ('Scapulothoracic', 'Posterior (+) / anterior (-) tilt'): 'positive = posterior tilt; negative = anterior tilt',
}

wb = Workbook()

# =====================================================================
# README
# =====================================================================
ws = wb.active; ws.title = 'README'
z = d['zone']['Glenohumeral']
lines = [
    ('Kinematics — SPM1D results (group and individual)', F_TITLE),
    ('Project STIM_KC — glenohumeral (GH) and scapulothoracic (ST) kinematics, 10 participants (P001–P010), 7 FES conditions, ANALYTIC2 task (coronal-plane arm elevation).', F_BASE),
    ('Source: cache_glenohumeral_all_comp.mat and cache_scapulothoracic_all_comp.mat, produced by extract_glenohumeral_kinematics_all_comp.m / extract_scapular_kinematics_all_comp.m (2-Final processing/Results). No value in this workbook was typed by hand.', F_NOTE),
    ('', F_BASE),
    ('Data', F_BOLD),
    ('• Joint angles (deg) from K-LAB Euler decomposition: GH = humerus relative to scapula (XZY sequence), ST = scapula relative to thorax (YXZ sequence). Each cycle is time-normalised to 101 points (0–100 % of the movement cycle).', F_BASE),
    ('• Per participant and condition: mean over cycles within each trial, then mean over the 3 blocks (trials) of that condition → one curve per participant × condition × DOF (sheet Curves_Individual).', F_BASE),
    ('', F_BASE),
    ('Statistics (SPM1D, Pataky 2010)', F_BOLD),
    ('• Group level (N = 10): one-way repeated-measures ANOVA across the 7 conditions per DOF — spm1d.stats.nonparam.anova1rm, non-parametric permutation test (Monte Carlo, 10 000 iterations, rng(0)), α = 0.05. With 10 000 iterations the smallest attainable p is 1e-4.', F_BASE),
    ('• Post-hoc (only for DOFs with a significant ANOVA — all 6 DOFs here): paired t-tests on all C(7,2) = 21 condition pairs (spm1d.stats.ttest_paired, two-tailed, Random Field Theory inference).', F_BASE),
    ('• Multiple-comparison correction: Holm-Bonferroni step-down (Holm 1979) over the 21 pairs of each DOF, family-wise α = 0.05. Each test gets one test-level p-value = RFT probability that the maximum of the 1D t field exceeds the observed max |t| (two-tailed). The k-th smallest p-value is compared with α/(m − k + 1), stopping at the first non-rejected test. Rejected tests are then thresholded at their own Holm α, which yields the reported significant windows (supra-threshold clusters) and cluster p-values. Sheet Group_PostHoc_Holm reproduces the Holm decision with live Excel formulas.', F_BASE),
    ('• Individual level (exploratory): same ANOVA + Holm-corrected post-hoc within each participant, using the 3 blocks as observations (N = 3, df = 2). Low power — exploratory only.', F_BASE),
    ('', F_BASE),
    ('Window angles', F_BOLD),
    ('• For each significant window, each participant\'s curve is averaged over the window samples; group values are mean ± SD across the 10 participants of these window-averaged angles; Diff = paired difference B − A (mean ± SD). Individual rows report the participant\'s own window-averaged values.', F_BASE),
    ('', F_BASE),
    ('Non-interpretable zone', F_BOLD),
    (f'• GH and ST angles are not interpreted where humerothoracic elevation exceeds {z["threshold"]:g}° (skin-marker scapular tracking becomes unreliable). With the group-mean humerothoracic elevation, this corresponds to {z["windows"][0]:.1f}–{z["windows"][1]:.1f} % of the cycle (sheet Exclusion_Zone). Statistics were still computed on the full cycle; the "% of window in zone" columns flag windows overlapping it.', F_BASE),
    ('', F_BASE),
    ('Sign conventions', F_BOLD),
]
r = 1
for text, font in lines:
    c = ws.cell(row=r, column=1, value=text); c.font = font
    c.alignment = Alignment(wrap_text=True, vertical='top')
    r += 1
ws.cell(row=r, column=1, value='Joint').font = F_HEAD; ws.cell(row=r, column=1).fill = FILL_HEAD
ws.cell(row=r, column=2, value='DOF').font = F_HEAD; ws.cell(row=r, column=2).fill = FILL_HEAD
ws.cell(row=r, column=3, value='Convention').font = F_HEAD; ws.cell(row=r, column=3).fill = FILL_HEAD
r += 1
for (j, dof), conv in CONV.items():
    for k, v in enumerate([j, dof, conv], 1):
        ws.cell(row=r, column=k, value=v).font = F_BASE
    r += 1
r += 1
ws.cell(row=r, column=1, value='Sheets').font = F_BOLD; r += 1
for name, desc in [
    ('Group_ANOVA', 'Group-level RM-ANOVA supra-threshold clusters (window, p) per joint × DOF.'),
    ('Group_PostHoc_Holm', 'All 126 group post-hoc tests (2 joints × 3 DOF × 21 pairs): test-level p, Holm rank/threshold/decision (Excel formulas) and SPM1D decision, with a consistency check.'),
    ('Group_Windows', 'Every significant group post-hoc window: % cycle, cluster p, window angles (mean ± SD), paired difference, overlap with the non-interpretable zone, participants also individually significant.'),
    ('Individual_Windows', 'Every significant individual post-hoc window (exploratory, N = 3 blocks), with the participant\'s window-averaged angles.'),
    ('Curves_Individual', 'Mean curve (101 points) of every participant × condition × DOF.'),
    ('Curves_Group', 'Group mean and SD curves (Excel formulas over Curves_Individual).'),
    ('Exclusion_Zone', 'Humerothoracic elevation group-mean curve and resulting non-interpretable window.'),
]:
    ws.cell(row=r, column=1, value=name).font = F_BOLD
    ws.cell(row=r, column=2, value=desc).font = F_BASE
    r += 1
r += 1
ws.cell(row=r, column=1, value='p-values below 0.001 are displayed as "<0.001" (cell format; the exact value is kept in the cell). Cluster p-values reported as 0 by SPM1D correspond to p < 1e-4.').font = F_NOTE
ws.column_dimensions['A'].width = 22; ws.column_dimensions['B'].width = 40; ws.column_dimensions['C'].width = 70
for rr, (text, font) in enumerate(lines, 1):
    ws.merge_cells(start_row=rr, start_column=1, end_row=rr, end_column=3)
    nl = max(1, math.ceil(len(text) / 125))
    ws.row_dimensions[rr].height = 20 if rr == 1 else 14 * nl + 2

# =====================================================================
# Exclusion_Zone
# =====================================================================
wz = wb.create_sheet('Exclusion_Zone')
title(wz, 'Non-interpretable zone (humerothoracic elevation > threshold)', 'Group-mean humerothoracic elevation (10 participants, all conditions pooled; + = elevation). Window edges linearly interpolated.')
header(wz, 4, ['Joint', 'Threshold (deg)', 'Zone start (% cycle)', 'Zone end (% cycle)', 'Peak HT elevation (deg)', 'Peak at (% cycle)'], [18, 14, 14, 14, 14, 14], freeze=False)
ZROW = {}
row = 5
for jn in ['Glenohumeral', 'Scapulothoracic']:
    zz = d['zone'][jn]
    wz.cell(row=row, column=1, value=jn)
    wz.cell(row=row, column=2, value=zz['threshold'])
    wz.cell(row=row, column=3, value=zz['windows'][0]).number_format = PCT
    wz.cell(row=row, column=4, value=zz['windows'][1]).number_format = PCT
    ZROW[jn] = row
    row += 1
style_body(wz, 5, row-1, 6)
crow = row + 2
wz.cell(row=crow, column=1, value='% cycle').font = F_HEAD; wz.cell(row=crow, column=1).fill = FILL_HEAD
wz.cell(row=crow, column=2, value='HT elevation (deg)').font = F_HEAD; wz.cell(row=crow, column=2).fill = FILL_HEAD
ht = d['zone']['Glenohumeral']['htGroupMean']
for i, v in enumerate(ht):
    wz.cell(row=crow+1+i, column=1, value=i).font = F_BASE
    c = wz.cell(row=crow+1+i, column=2, value=val(v)); c.font = F_BASE; c.number_format = ANG
    if v > d['zone']['Glenohumeral']['threshold']:
        c.fill = FILL_ZONE
r1, r2 = crow+1, crow+101
for jn, rr in ZROW.items():
    wz.cell(row=rr, column=5, value=f'=MAX($B${r1}:$B${r2})').number_format = ANG
    wz.cell(row=rr, column=6, value=f'=INDEX($A${r1}:$A${r2},MATCH(E{rr},$B${r1}:$B${r2},0))')
    for cc in (5, 6): wz.cell(row=rr, column=cc).font = F_BASE
wz.cell(row=crow-1, column=1, value='Same humerothoracic signal for both joints (zone identical).').font = F_NOTE
ZS = lambda jcell: f"INDEX(Exclusion_Zone!$C$5:$C$6,MATCH({jcell},Exclusion_Zone!$A$5:$A$6,0))"
ZE = lambda jcell: f"INDEX(Exclusion_Zone!$D$5:$D$6,MATCH({jcell},Exclusion_Zone!$A$5:$A$6,0))"

def overlap_formula(jc, sc, ec):
    zs, ze = ZS(jc), ZE(jc)
    return (f'=IF({ec}-{sc}<=0,IF(AND({sc}>={zs},{sc}<={ze}),100,0),'
            f'100*MAX(0,MIN({ec},{ze})-MAX({sc},{zs}))/({ec}-{sc}))')

# =====================================================================
# Group_ANOVA
# =====================================================================
wa = wb.create_sheet('Group_ANOVA')
title(wa, 'Group-level repeated-measures ANOVA (7 conditions, N = 10)', 'spm1d.stats.nonparam.anova1rm — permutation (10 000 iterations), α = 0.05. One row per supra-threshold cluster.')
cols = ['Joint', 'DOF', 'Significant', 'Window start (% cycle)', 'Window end (% cycle)', 'Duration (% cycle)', 'p (permutation)']
header(wa, 4, cols, [16, 36, 12, 14, 14, 13, 14])
row = 5
for a in d['anova']:
    wa.cell(row=row, column=1, value=a['joint']); wa.cell(row=row, column=2, value=a['dof'])
    wa.cell(row=row, column=3, value='Yes' if a['sig'] else 'No')
    if a['sig']:
        wa.cell(row=row, column=4, value=val(a['start'])).number_format = PCT
        wa.cell(row=row, column=5, value=val(a['end'])).number_format = PCT
        wa.cell(row=row, column=6, value=f'=E{row}-D{row}').number_format = PCT
        wa.cell(row=row, column=7, value=val(a['p'])).number_format = P_FMT
    row += 1
style_body(wa, 5, row-1, len(cols))
wa.cell(row=row+1, column=1, value='p = 1e-4 is the resolution limit of 10 000 permutations (displayed "<0.001").').font = F_NOTE

# =====================================================================
# Group_PostHoc_Holm
# =====================================================================
wh = wb.create_sheet('Group_PostHoc_Holm')
title(wh, 'Group post-hoc — all pairwise tests and Holm-Bonferroni decision (N = 10)',
      'Family = 21 pairs of one joint × DOF. Columns F–J are live formulas reproducing the Holm step-down from the test-level p-values (column E); column K is the decision returned by the MATLAB pipeline; column L checks they agree.')
cols = ['Joint', 'DOF', 'Condition A', 'Condition B', 'Test-level p (RFT, max |t|)', 'Family size m', 'Rank k', 'Holm threshold α/(m−k+1)',
        'p − threshold', 'Holm decision (Excel)', 'SPM1D decision (MATLAB)', 'Check', 'Participants also individually significant (n)', 'Participant IDs']
header(wh, 4, cols, [16, 34, 13, 13, 15, 10, 9, 15, 13, 14, 14, 9, 15, 26])
ph = d['posthoc']
# une ligne par test (pas par cluster)
seen = set(); tests = []
for p in ph:
    key = (p['joint'], p['dof'], p['condA'], p['condB'])
    if key in seen: continue
    seen.add(key); tests.append(p)
r1 = 5; r2 = r1 + len(tests) - 1
ALPHA_CELL = 'Group_PostHoc_Holm!$B$3'
wh['A3'] = 'Family-wise α'; wh['A3'].font = F_BOLD
wh['B3'] = 0.05; wh['B3'].font = Font(name=FONT, size=10, color='0000FF')
wh['B3'].comment = Comment('Family-wise alpha used for the Holm-Bonferroni correction (study design value). Changing it re-evaluates columns H–J (not the MATLAB decision in K).', 'pipeline')
for i, t in enumerate(tests):
    rr = r1 + i
    rng = lambda col: f'${col}${r1}:${col}${r2}'
    wh.cell(row=rr, column=1, value=t['joint']); wh.cell(row=rr, column=2, value=t['dof'])
    wh.cell(row=rr, column=3, value=t['condA']); wh.cell(row=rr, column=4, value=t['condB'])
    wh.cell(row=rr, column=5, value=val(t['p_test'])).number_format = P_FMT4
    wh.cell(row=rr, column=6, value=f'=COUNTIFS({rng("A")},A{rr},{rng("B")},B{rr})')
    wh.cell(row=rr, column=7, value=f'=COUNTIFS({rng("A")},A{rr},{rng("B")},B{rr},{rng("E")},"<"&E{rr})+1')
    wh.cell(row=rr, column=8, value=f'=$B$3/(F{rr}-G{rr}+1)').number_format = '0.00000'
    wh.cell(row=rr, column=9, value=f'=E{rr}-H{rr}').number_format = '0.00000'
    wh.cell(row=rr, column=10, value=f'=IF(_xlfn.MAXIFS({rng("I")},{rng("A")},A{rr},{rng("B")},B{rr},{rng("E")},"<="&E{rr})<0,"Significant","n.s.")')
    wh.cell(row=rr, column=11, value='Significant' if t['sig'] else 'n.s.')
    wh.cell(row=rr, column=12, value=f'=IF(J{rr}=K{rr},"OK","CHECK")')
    wh.cell(row=rr, column=13, value=t['nPatSig'])
    wh.cell(row=rr, column=14, value=t['patSig'] or None)
style_body(wh, r1, r2, len(cols))
for i, t in enumerate(tests):
    if t['sig']:
        for c in range(1, len(cols)+1):
            wh.cell(row=r1+i, column=c).fill = FILL_SIG
wh.cell(row=r2+2, column=1, value='Test-level p = 2 × RFT survival probability of the observed max |t| (spm1d.rft1d.t.sf_resels, two-tailed). p < threshold at rank k and at every lower rank ⇒ rejected (Holm step-down). Green rows = significant.').font = F_NOTE

# =====================================================================
# Group_Windows
# =====================================================================
wg = wb.create_sheet('Group_Windows')
title(wg, 'Group post-hoc — significant windows and window angles (N = 10)',
      'Angles in degrees (see README for sign conventions). Mean ± SD across participants of each participant\'s window-averaged angle; Diff = paired B − A.')
cols = ['Joint', 'DOF', 'Condition A', 'Condition B', 'Window start (% cycle)', 'Window end (% cycle)', 'Duration (% cycle)', 'Cluster p',
        'Holm threshold', 'Angle A mean (deg)', 'Angle A SD', 'Angle B mean (deg)', 'Angle B SD', 'Diff B − A mean (deg)', 'Diff B − A SD',
        '% of window in non-interpretable zone', 'Participants also individually significant (n)', 'Participant IDs', 'Sign convention']
header(wg, 4, cols, [16, 34, 13, 13, 12, 12, 11, 11, 11, 12, 10, 12, 10, 12, 10, 14, 14, 24, 52])
row = 5
for p in [x for x in ph if x['sig']]:
    vals = [p['joint'], p['dof'], p['condA'], p['condB'], p['start'], p['end'], None, p['p_cluster'], p['alpha_holm'],
            p['meanA'], p['sdA'], p['meanB'], p['sdB'], p['diff'], p['sdDiff'], None, p['nPatSig'], p['patSig'] or None,
            CONV[(p['joint'], p['dof'])]]
    for c, v in enumerate(vals, 1):
        wg.cell(row=row, column=c, value=val(v))
    wg.cell(row=row, column=7, value=f'=F{row}-E{row}')
    wg.cell(row=row, column=16, value=overlap_formula(f'A{row}', f'E{row}', f'F{row}'))
    for c in (5, 6, 7, 16): wg.cell(row=row, column=c).number_format = PCT
    wg.cell(row=row, column=8).number_format = P_FMT
    wg.cell(row=row, column=9).number_format = '0.00000'
    for c in range(10, 16): wg.cell(row=row, column=c).number_format = ANG
    row += 1
style_body(wg, 5, row-1, len(cols))
wg.cell(row=row+1, column=1, value='Window samples: nodes round(start)…round(end) of the 101-point cycle (same definition as the article tables). Cluster p computed by SPM1D (RFT) at the test\'s Holm threshold.').font = F_NOTE

# =====================================================================
# Individual_Windows
# =====================================================================
wi = wb.create_sheet('Individual_Windows')
title(wi, 'Individual post-hoc — significant windows (exploratory, N = 3 blocks per participant)',
      'Within-participant RM-ANOVA + Holm-corrected paired t-tests on the 3 blocks (df = 2). Angles = participant\'s block-averaged curve, averaged over the window.')
cols = ['Participant', 'Joint', 'DOF', 'Condition A', 'Condition B', 'Window start (% cycle)', 'Window end (% cycle)', 'Duration (% cycle)',
        'Cluster p', 'Angle A (deg)', 'Angle B (deg)', 'Diff B − A (deg)', '% of window in non-interpretable zone', 'Also significant at group level']
header(wi, 4, cols, [11, 16, 34, 13, 13, 12, 12, 11, 11, 11, 11, 11, 14, 13])
grpSig = {(p['joint'], p['dof'], p['condA'], p['condB']) for p in ph if p['sig']}
ind = sorted(d['indiv'], key=lambda x: (x['patient'], x['joint'], x['dof'], x['start']))
row = 5
for p in ind:
    vals = [p['patient'], p['joint'], p['dof'], p['condA'], p['condB'], p['start'], p['end'], None, p['p_cluster'], p['meanA'], p['meanB']]
    for c, v in enumerate(vals, 1):
        wi.cell(row=row, column=c, value=val(v))
    wi.cell(row=row, column=8, value=f'=G{row}-F{row}')
    wi.cell(row=row, column=12, value=f'=K{row}-J{row}')
    wi.cell(row=row, column=13, value=overlap_formula(f'B{row}', f'F{row}', f'G{row}'))
    wi.cell(row=row, column=14, value='Yes' if (p['joint'], p['dof'], p['condA'], p['condB']) in grpSig else 'No')
    for c in (6, 7, 8, 13): wi.cell(row=row, column=c).number_format = PCT
    wi.cell(row=row, column=9).number_format = P_FMT
    for c in (10, 11, 12): wi.cell(row=row, column=c).number_format = ANG
    row += 1
style_body(wi, 5, row-1, len(cols))
wi.cell(row=row+1, column=1, value='Exploratory: N = 3 blocks per participant gives very low degrees of freedom; these results support, but do not replace, the group analysis.').font = F_NOTE

# =====================================================================
# Curves_Individual
# =====================================================================
wc = wb.create_sheet('Curves_Individual')
title(wc, 'Individual mean curves (deg), 101 points per cycle', 'One row per participant × condition × DOF (mean over cycles, then over the 3 blocks of the condition). Columns = % of the movement cycle.')
cols = ['Joint', 'DOF', 'Condition', 'Participant'] + [str(i) for i in range(101)]
header(wc, 4, cols, [16, 34, 13, 11] + [6.5]*101)
row = 5
blocks = []  # (joint, dof, cond, firstRow, lastRow)
cur_key = None
for c in d['curves']:
    key = (c['joint'], c['dof'], c['cond'])
    if key != cur_key:
        if cur_key: blocks.append((*cur_key, first, row-1))
        cur_key = key; first = row
    wc.cell(row=row, column=1, value=c['joint']); wc.cell(row=row, column=2, value=c['dof'])
    wc.cell(row=row, column=3, value=c['cond']); wc.cell(row=row, column=4, value=c['patient'])
    for i, v in enumerate(c['y']):
        cell = wc.cell(row=row, column=5+i, value=val(v)); cell.number_format = ANG; cell.font = F_BASE
    for k in range(1, 5): wc.cell(row=row, column=k).font = F_BASE
    row += 1
blocks.append((*cur_key, first, row-1))
wc.freeze_panes = 'E5'

# =====================================================================
# Curves_Group
# =====================================================================
wq = wb.create_sheet('Curves_Group')
title(wq, 'Group mean and SD curves (deg), N = 10', 'Live formulas: AVERAGE / STDEV over the 10 participant rows of Curves_Individual.')
cols = ['Joint', 'DOF', 'Condition', 'Statistic'] + [str(i) for i in range(101)]
header(wq, 4, cols, [16, 34, 13, 10] + [6.5]*101)
row = 5
for (jn, dof, cond, fr, lr) in blocks:
    for stat, fn in (('Mean', 'AVERAGE'), ('SD', 'STDEV')):
        for k, v in enumerate([jn, dof, cond, stat], 1):
            wq.cell(row=row, column=k, value=v).font = F_BASE
        for i in range(101):
            col = get_column_letter(5+i)
            cell = wq.cell(row=row, column=5+i, value=f'={fn}(Curves_Individual!{col}{fr}:{col}{lr})')
            cell.number_format = ANG; cell.font = F_BASE
        row += 1
wq.freeze_panes = 'E5'

for wsx in wb.worksheets:
    wsx.sheet_view.zoomScale = 100
wb.calculation.fullCalcOnLoad = True  # Excel recalcule toutes les formules a l'ouverture
wb.save(OUT)
print('saved', OUT, 'tests', len(tests), 'groupWin', sum(1 for x in ph if x['sig']), 'indiv', len(ind), 'blocks', len(blocks))
